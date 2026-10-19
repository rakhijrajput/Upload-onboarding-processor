"""Read CSV and XLSX uploads into normalized row dictionaries."""
import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLUMNS = ("employee_id", "name", "department", "role", "joining_date", "manager")


def normalize_header(value: Any) -> str:
    return "_".join(str(value or "").strip().lower().replace("-", " ").split())


def normalize_row(raw: dict[str, Any]) -> dict[str, Any]:
    return {normalize_header(key): value for key, value in raw.items() if key is not None}


def read_employee_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV must use UTF-8 encoding.") from exc
        reader = csv.DictReader(StringIO(text))
        if not reader.fieldnames:
            raise ValueError("CSV has no header row.")
        return [normalize_row(row) for row in reader]

    if suffix == ".xlsx":
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("Could not read this XLSX file. Check that it is a valid Excel workbook.") from exc
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("XLSX worksheet has no header row.")
            normalized_headers = [normalize_header(value) for value in headers]
            rows = []
            for values_row in values:
                if not any(value is not None and str(value).strip() for value in values_row):
                    continue
                rows.append({header: value for header, value in zip(normalized_headers, values_row) if header})
            return rows
        finally:
            workbook.close()

    raise ValueError("Unsupported input type. Upload a .csv or .xlsx file.")
